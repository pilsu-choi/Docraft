import csv
import io
import shutil
from pathlib import Path

from docx import Document
from openpyxl import load_workbook
from PIL import Image
from pypdf import PdfReader


class ParseError(ValueError):
    pass


def block(text, kind="text", page=None, bbox=None, **extra):
    return {"type": kind, "page": page, "bbox": bbox, "text": str(text), **extra}


def parse_pdf(path):
    result = []
    for number, page in enumerate(PdfReader(path).pages, 1):
        width, height = float(page.mediabox.width), float(page.mediabox.height)

        def visit(text, _cm, tm, _font, _size):
            value = text.strip()
            if value:
                x, y = float(tm[4]), float(tm[5])
                # pypdf exposes the text origin, not glyph extents; keep point bbox honest.
                result.append(block(value, page=number, bbox=[x, max(0, height-y), x, max(0, height-y)], page_size=[width, height]))

        page.extract_text(visitor_text=visit)
    if not result:
        raise ParseError("PDF에 추출 가능한 텍스트가 없습니다. 스캔 PDF는 OCR 제공자 설정이 필요합니다.")
    return result


def parse_image(path):
    if not shutil.which("tesseract"):
        raise ParseError("이미지 OCR을 위해 시스템에 Tesseract를 설치하거나 AI 제공자를 설정해 주세요.")
    import pytesseract

    image = Image.open(path)
    data = pytesseract.image_to_data(image, output_type=pytesseract.Output.DICT)
    result = []
    for i, value in enumerate(data["text"]):
        value = value.strip()
        if value:
            x, y, w, h = (data[k][i] for k in ("left", "top", "width", "height"))
            result.append(block(value, page=1, bbox=[x, y, x+w, y+h], confidence=float(data["conf"][i]) / 100))
    if not result:
        raise ParseError("이미지에서 텍스트를 찾지 못했습니다.")
    return result


def parse_docx(path):
    doc = Document(path)
    result = [block(p.text, "heading" if p.style.name.startswith("Heading") else "text") for p in doc.paragraphs if p.text.strip()]
    for table_no, table in enumerate(doc.tables):
        rows = [[cell.text.strip() for cell in row.cells] for row in table.rows]
        result.append(block("\n".join(" | ".join(row) for row in rows), "table", rows=rows, table=table_no))
    return result


def parse_xlsx(path):
    result = []
    book = load_workbook(path, read_only=True, data_only=True)
    for sheet in book.worksheets:
        rows = [["" if value is None else str(value) for value in row] for row in sheet.iter_rows(values_only=True)]
        rows = [row for row in rows if any(row)]
        if rows:
            result.append(block("\n".join(" | ".join(row) for row in rows), "table", rows=rows, sheet=sheet.title))
    return result


def parse_csv(path):
    raw = Path(path).read_bytes()
    text = raw.decode("utf-8-sig", errors="replace")
    rows = list(csv.reader(io.StringIO(text)))
    return [block("\n".join(" | ".join(row) for row in rows), "table", rows=rows)]


def parse_text(path):
    text = Path(path).read_text(encoding="utf-8", errors="replace")
    return [block(line) for line in text.splitlines() if line.strip()]


def parse(path, filename, media_type):
    suffix = Path(filename).suffix.lower()
    if suffix == ".pdf":
        blocks = parse_pdf(path)
    elif suffix in {".png", ".jpg", ".jpeg", ".tif", ".tiff", ".webp"}:
        blocks = parse_image(path)
    elif suffix == ".docx":
        blocks = parse_docx(path)
    elif suffix == ".xlsx":
        blocks = parse_xlsx(path)
    elif suffix == ".csv":
        blocks = parse_csv(path)
    elif suffix in {".txt", ".md"} or media_type.startswith("text/"):
        blocks = parse_text(path)
    else:
        raise ParseError(f"지원하지 않는 파일 형식입니다: {suffix or media_type}")
    if not blocks:
        raise ParseError("문서에서 내용을 찾지 못했습니다.")
    markdown = "\n\n".join(b["text"] if b["type"] != "table" else f"```text\n{b['text']}\n```" for b in blocks)
    return markdown, blocks
