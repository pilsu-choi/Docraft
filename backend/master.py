"""KCD 상병·EDI 수가 마스터 사전 조회.

원본 행 (family, code, name)을 세 소스 중 먼저 되는 것에서 얻어(``_rows``) 지연 적재하고,
코드가 마스터에 있는 행에 한해 명칭의 1글자 OCR 오인식(편축→편측, 부문→부분)을 되돌린다.
명칭에서 코드를 역추론하지는 않는다 — 코드가 틀렸을 때 엉뚱한 명칭을 끼워 넣는 쪽이
미검출보다 나쁘다(harness-v2 Arbitration 원칙).

1. ``HARNESS_DATABASE_URL`` — harness-v2 Postgres의 ``code_entry``/``code_system``을 읽기전용 재사용.
2. Docraft 자체 DB ``master_code`` — 이미 적재돼 있으면 그대로 쓴다.
3. ``MASTER_SOURCE_DIR`` — harness 원본 파일(KCD_CODE_*.csv, 수가코드_*.xlsx, 약가_*.tar.gz,
   치료재료_전체_*.tar.gz)을 파싱해 ``master_code``에 적재한다(advisory lock으로 동시 기동 보호).

아무 소스도 없으면 조용히 비활성. 새 고시 원장이 오면 ``python -m backend.master --source <dir>``로
강제 재적재한다.
"""

import csv
import functools
import gzip
import logging
import os
import re
import tarfile
from pathlib import Path

import psycopg

from backend import db

logger = logging.getLogger(__name__)

MAX_EDITS = 1  # 마스터 명칭과 어긋나도 되는 글자 수. 76건 평가에서 1이 최선(2는 순이득 -1)
EDI_MIN = 5  # 뒤에서 절단해 다시 찾을 때 남겨야 할 최소 길이(harness E4)
DRUG_LEN = 9  # 9자리 약가코드는 절단하지 않는다 — 다른 약에 잘못 붙는다

_STRIP = re.compile(r"[\s\-_.]+")  # 코드 표기 차이
_LOOSE = re.compile(r"[\s\-_,()\[\]{}]")  # 명칭 표기 차이(고시는 하이픈, 인쇄물은 괄호)


def code(text) -> str | None:
    """조회용 코드 표기: 공백·하이픈·밑줄·점(KCD의 ``M81.99``)을 빼고 대문자로."""
    if not text:
        return None
    return _STRIP.sub("", str(text)).upper() or None


def _system(family) -> str | None:
    """harness family(``KCD``·``EDI``·``EDI:약가`` 등)를 Docraft system(``kcd``·``edi``)으로."""
    family = (family or "").upper()
    if family.startswith("KCD"):
        return "kcd"
    if family.startswith("EDI"):
        return "edi"
    return None


# ── 원본 소스 ────────────────────────────────────────────────────────────────

def _from_harness() -> list[tuple[str, str, str]] | None:
    """harness-v2 Postgres를 읽기전용으로 재사용."""
    url = os.getenv("HARNESS_DATABASE_URL", "").strip()
    if not url:
        return None
    try:
        with psycopg.connect(url, connect_timeout=5) as conn:
            rows = conn.execute("""
                SELECT DISTINCT s.family, e.code, e.name
                FROM code_entry e JOIN code_system s USING (system_id)
                WHERE s.enabled AND s.family IN ('KCD', 'EDI')
            """).fetchall()
    except psycopg.Error as exc:
        logger.warning("master: harness DB 조회 실패(%s), 다음 소스로 넘어간다", exc)
        return None
    return [tuple(row) for row in rows] or None


def _from_docraft_db() -> list[tuple[str, str, str]] | None:
    """Docraft 자체 DB의 ``master_code``(이미 적재돼 있으면)."""
    try:
        with db.connect() as conn:
            rows = conn.execute("SELECT family,code,name FROM master_code").fetchall()
    except psycopg.Error as exc:
        logger.warning("master: docraft DB 조회 실패(%s), 다음 소스로 넘어간다", exc)
        return None
    return [(row["family"], row["code"], row["name"]) for row in rows] or None


def _latest(root: Path, pattern: str) -> Path | None:
    """여러 판이 있으면 최신 파일명을 쓴다(고시 갱신 시 파일이 늘어난다)."""
    matches = sorted(root.glob(pattern))
    return matches[-1] if matches else None


def _kcd_rows(root: Path):
    path = _latest(root, "KCD_CODE_*.csv")
    if not path:
        return
    with path.open(encoding="cp949", newline="") as handle:
        for row in csv.DictReader(handle):
            yield "KCD", row.get("상병기호"), row.get("한글명")


def _edi_rows(root: Path):
    path = _latest(root, "수가코드_*.xlsx")
    if not path:
        return
    from openpyxl import load_workbook

    book = load_workbook(path, read_only=True, data_only=True)
    try:
        for sheet in book.worksheets:
            header = None
            for cells in sheet.iter_rows(values_only=True):
                names = [str(cell).strip() if cell is not None else "" for cell in cells]
                if header is None:
                    if "수가코드" in names:  # 시트마다 머리글 행 위치·열 순서가 다르다
                        header = {"code": names.index("수가코드"), "name": names.index("한글명")}
                    continue
                yield "EDI", cells[header["code"]], cells[header["name"]]
    finally:
        book.close()


def _tar_rows(root: Path, family: str, pattern: str, member: str):
    path = _latest(root, pattern)
    if not path:
        return
    with tarfile.open(path) as tar:
        entry = next((m for m in tar.getmembers() if member in m.name), None)
        if entry is None:
            return
        raw = tar.extractfile(entry).read()
        if entry.name.endswith(".gz"):
            raw = gzip.decompress(raw)
    for row in csv.DictReader(raw.decode("utf-8-sig").splitlines()):
        yield family, row.get("code"), row.get("name")


def _parse_source(root: Path) -> list[tuple[str, str, str]]:
    """harness 마스터 원본 디렉터리를 (family, code, name) 행으로 푼다. 단가는 쓰지 않는다."""
    seen: set[tuple[str, str, str]] = set()
    rows: list[tuple[str, str, str]] = []
    for family, raw_code, raw_name in (
        *_kcd_rows(root), *_edi_rows(root),
        *_tar_rows(root, "EDI:약가", "약가_*.tar.gz", "dim_drug.csv"),
        *_tar_rows(root, "EDI:치료재료", "치료재료_전체_*.tar.gz", "dim_material.csv"),
    ):
        entry = ((raw_code or "").strip(), (raw_name or "").strip())
        if not entry[0] or not entry[1] or (family, *entry) in seen:
            continue
        seen.add((family, *entry))
        rows.append((family, *entry))
    return rows


def _replace(conn, rows) -> None:
    """``master_code``를 COPY로 통째로 교체한다. 호출부가 advisory lock 구간 안에서 부른다."""
    conn.execute("TRUNCATE master_code")
    with conn.connection.cursor() as cur, cur.copy("COPY master_code (family,code,name) FROM STDIN") as copy:
        for row in rows:
            copy.write_row(row)


def _load_source() -> list[tuple[str, str, str]] | None:
    """``MASTER_SOURCE_DIR``의 원본을 파싱해 비어 있으면 ``master_code``에 적재한다."""
    source = os.getenv("MASTER_SOURCE_DIR", "").strip()
    if not source or not Path(source).is_dir():
        return None
    rows = _parse_source(Path(source))
    if not rows:
        return None
    with db.connect() as conn:
        conn.execute("SELECT pg_advisory_xact_lock(hashtext('docraft.master_load'))")  # API·worker 동시 기동 보호
        if not conn.execute("SELECT 1 FROM master_code LIMIT 1").fetchone():
            _replace(conn, rows)
    return rows


def reload_from(source: Path) -> int:
    """새 고시 원장이 왔을 때 원본에서 ``master_code``를 강제로 다시 채운다."""
    rows = _parse_source(source)
    with db.connect() as conn:
        conn.execute("SELECT pg_advisory_xact_lock(hashtext('docraft.master_load'))")
        _replace(conn, rows)
    _tables.cache_clear()
    return len(rows)


def _rows() -> list[tuple[str, str, str]]:
    for loader, label in ((_from_harness, "harness DB 재사용"), (_from_docraft_db, "docraft DB 재사용"),
                          (_load_source, "원본에서 신규 적재")):
        rows = loader()
        if rows:
            logger.info("master: %s (%d건)", label, len(rows))
            return rows
    return []


def _build(rows) -> dict[str, dict[str, list[str]]]:
    tables: dict[str, dict[str, list[str]]] = {}
    for family, raw_code, raw_name in rows:
        system = _system(family)
        key = code(raw_code)
        name = (raw_name or "").strip()
        if not system or not key or not name:
            continue
        table = tables.setdefault(system, {})
        if name not in table.setdefault(key, []):
            table[key].append(name)
    for system, table in tables.items():
        logger.info("master: %s 코드 %d건 적재", system, len(table))
    return tables


@functools.lru_cache(maxsize=1)
def _tables() -> dict[str, dict[str, list[str]]]:
    return _build(_rows())


def names(system: str, value) -> list[str]:
    """코드에 딸린 마스터 명칭들. EDI는 정확 일치가 없으면 뒤에서 한 자씩 줄여 다시 찾는다(harness E3·E4)."""
    table = _tables().get(system)
    key = code(value)
    if not table or not key:
        return []
    if key in table:
        return table[key]
    if system != "edi" or (len(key) == DRUG_LEN and key.isdigit()):
        return []
    for end in range(len(key) - 1, EDI_MIN - 1, -1):
        if key[:end] in table:
            return table[key[:end]]
    return []


def _letters(text: str) -> tuple[str, list[int]]:
    """비교용 글자열과 각 글자의 원문 위치. 띄어쓰기·괄호·하이픈은 표기 차이라 빼고 센다."""
    kept = [(index, char) for index, char in enumerate(text) if not _LOOSE.match(char)]
    return "".join(char.lower() for _, char in kept), [index for index, _ in kept]


def correct_name(system: str, value, name) -> str | None:
    """코드가 마스터에 있고 명칭이 마스터 명칭과 ``MAX_EDITS`` 글자만 다를 때, 그 글자만 고친 명칭.

    마스터 명칭을 통째로 쓰지는 않는다. 라벨도 정답도 문서에 인쇄된 표기인데 마스터는 고시 표기라
    띄어쓰기·괄호·어순이 달라, 통째로 갈아끼우면 맞던 칸까지 어긋난다(76건 평가에서 순손실 -13).
    글자 수가 같을 때만 보므로 낱말이 더 있거나 빠진 경우는 교정하지 않는다.
    """
    text = str(name or "").strip()
    given, positions = _letters(text)
    if not given:
        return None
    for candidate in names(system, value):
        other, other_positions = _letters(candidate)
        if len(other) != len(given):
            continue
        diff = [index for index, (a, b) in enumerate(zip(given, other)) if a != b]
        if not diff or len(diff) > MAX_EDITS:
            continue
        chars = list(text)
        for index in diff:
            chars[positions[index]] = candidate[other_positions[index]]
        return "".join(chars)
    return None


def ready() -> bool:
    return bool(_tables())


if __name__ == "__main__":
    import argparse

    parser = argparse.ArgumentParser(description="KCD·EDI 마스터를 원본에서 docraft DB로 강제 재적재한다.")
    parser.add_argument("--source", type=Path, required=True, help="harness 마스터 원본 디렉터리")
    args = parser.parse_args()
    if not args.source.is_dir():
        parser.error(f"원본 디렉터리가 없습니다: {args.source}")
    db.init_db()
    count = reload_from(args.source)
    print(f"master_code: {count:,}건 재적재")
