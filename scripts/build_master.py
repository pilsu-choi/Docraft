#!/usr/bin/env python3
"""harness-v2 마스터 원본을 ``data/master/*.csv``(code,name,unit_price)로 압축한다.

원본은 CP949 CSV·43MB 엑셀·tar.gz라 그대로 조회하기 무겁다. 여기서 코드·명칭·단가만 남긴
UTF-8 CSV로 줄여 두고, ``backend/master.py``가 그 파일만 읽는다. ``data/``는 gitignore 대상이라
결과물은 커밋되지 않으므로 이 스크립트를 한 번 돌려 만든다.

사용: python scripts/build_master.py [--source <harness-v2/docs/requirements/latest>] [--out data/master]
"""

import argparse
import csv
import gzip
import sys
import tarfile
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))

SOURCE = Path("/home/pilsu/projects/mirae-assets/harness-v2/docs/requirements/latest")
KCD_FILE = "KCD_CODE_20250930.csv"
EDI_FILE = "수가코드_250101_전체판.xlsx"


def _write(path: Path, rows) -> int:
    path.parent.mkdir(parents=True, exist_ok=True)
    seen, count = set(), 0
    with path.open("w", encoding="utf-8", newline="") as handle:
        writer = csv.writer(handle)
        writer.writerow(("code", "name", "unit_price"))
        for code, name, price in rows:
            code, name = (code or "").strip(), (name or "").strip()
            if not code or not name or (code, name) in seen:
                continue
            seen.add((code, name))
            writer.writerow((code, name, price or ""))
            count += 1
    return count


def _kcd(source: Path):
    with (source / KCD_FILE).open(encoding="cp949", newline="") as handle:
        for row in csv.DictReader(handle):
            yield row.get("상병기호"), row.get("한글명"), None


def _edi(source: Path):
    from openpyxl import load_workbook

    book = load_workbook(source / EDI_FILE, read_only=True, data_only=True)
    for sheet in book.worksheets:
        header = None
        for cells in sheet.iter_rows(values_only=True):
            names = [str(cell).strip() if cell is not None else "" for cell in cells]
            if header is None:
                if "수가코드" in names:  # 시트마다 머리글 행 위치·열 순서가 다르다
                    header = {"code": names.index("수가코드"), "name": names.index("한글명"),
                              "price": next((i for i, n in enumerate(names) if "단가" in n), None)}
                continue
            price = cells[header["price"]] if header["price"] is not None else None
            yield cells[header["code"]], cells[header["name"]], price
    book.close()


def _tar_csv(source: Path, pattern: str, member: str):
    """tar.gz 안의 CSV(gz일 수도 있다)를 DictReader로 흘린다. 여러 판이 있으면 최신 파일명을 쓴다."""
    archives = sorted(source.glob(pattern))
    if not archives:
        return
    with tarfile.open(archives[-1]) as tar:
        entry = next((m for m in tar.getmembers() if member in m.name), None)
        if entry is None:
            return
        raw = tar.extractfile(entry).read()
        if entry.name.endswith(".gz"):
            raw = gzip.decompress(raw)
    for row in csv.DictReader(raw.decode("utf-8-sig").splitlines()):
        yield row.get("code"), row.get("name"), row.get("상한금액")


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--source", type=Path, default=SOURCE, help="harness-v2 마스터 원본 디렉터리")
    parser.add_argument("--out", type=Path, default=ROOT / "data" / "master", help="조회 CSV를 쓸 디렉터리")
    args = parser.parse_args()
    if not args.source.is_dir():
        parser.error(f"원본 디렉터리가 없습니다: {args.source}")

    tasks = (("kcd.csv", _kcd(args.source)),
             ("edi.csv", _edi(args.source)),
             ("drug.csv", _tar_csv(args.source, "약가_*.tar.gz", "dim_drug.csv")),
             ("material.csv", _tar_csv(args.source, "치료재료_전체_*.tar.gz", "dim_material.csv")))
    for name, rows in tasks:
        print(f"{name}: {_write(args.out / name, rows):,}건", flush=True)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
